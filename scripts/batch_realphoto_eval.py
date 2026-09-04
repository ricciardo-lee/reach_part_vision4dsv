#!/usr/bin/env python3
"""Batch-evaluate real part photos against the local recognition API and write Excel."""

from __future__ import annotations

import argparse
import json
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Optional, Sequence, Set

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from realphoto_eval_lib import (  # noqa: E402
    DEFAULT_MAX_IMAGES,
    DEFAULT_PER_PREFIX,
    DEFAULT_SEED,
    MaterialGroup,
    UploadSet,
    eligible_groups,
    index_photos,
    prefix_counts,
    sample_by_prefix,
    select_upload_images,
)

DEFAULT_BASE_URL = "http://127.0.0.1:8000"
DEFAULT_DEVICE_ID = "realphoto-eval-batch"
RESULT_SHEET = "识别结果"
SAMPLE_SHEET = "抽样清单"
INFO_SHEET = "运行信息"
HEADER_FONT = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")

RESULT_HEADERS = [
    "品类前缀",
    "物料代码",
    "上传张数",
    "角度编号",
    "图片文件名",
    "图片路径",
    "public_id",
    "状态",
    "是否汽配",
    "识别品类",
    "映射品类",
    "映射品类代码",
    "置信度",
    "识别依据",
    "候选",
    "错误信息",
    "耗时秒",
    "测试时间",
    "人工判定",
    "备注",
]

SAMPLE_HEADERS = [
    "品类前缀",
    "物料代码",
    "可用张数",
    "可用角度",
    "将上传张数",
    "将上传角度",
    "是否已测",
]


class EvalApiError(RuntimeError):
    """Recognition API returned an error response."""


def _unwrap(payload: Any) -> Any:
    if not isinstance(payload, dict):
        return payload
    if payload.get("success") is False:
        message = payload.get("message") or payload.get("code") or "API 调用失败"
        raise EvalApiError(str(message))
    if "data" in payload:
        return payload["data"]
    return payload


class RecognizeClient:
    def __init__(
        self,
        base_url: str,
        *,
        timeout: float = 120.0,
        poll_interval: float = 3.0,
        poll_timeout: float = 180.0,
        device_id: str = DEFAULT_DEVICE_ID,
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self.poll_interval = poll_interval
        self.poll_timeout = poll_timeout
        self.device_id = device_id
        self._token: Optional[str] = None
        self._client = httpx.Client(base_url=self.base_url, timeout=timeout)

    def close(self) -> None:
        self._client.close()

    def __enter__(self) -> "RecognizeClient":
        return self

    def __exit__(self, *args: Any) -> None:
        self.close()

    def _headers(self) -> Dict[str, str]:
        if not self._token:
            return {}
        return {"X-Device-Token": self._token}

    def _request(self, method: str, path: str, **kwargs: Any) -> Any:
        delay = 2.0
        last_error: Optional[Exception] = None
        for attempt in range(8):
            try:
                response = self._client.request(method, path, headers=self._headers(), **kwargs)
            except httpx.TransportError as exc:
                last_error = exc
                time.sleep(delay)
                delay = min(delay * 2, 30)
                continue
            if response.status_code == 429:
                last_error = EvalApiError("请求过于频繁")
                time.sleep(delay)
                delay = min(delay * 2, 30)
                continue
            if response.status_code >= 500 and attempt < 4:
                last_error = EvalApiError(f"HTTP {response.status_code}: {response.text[:200]}")
                time.sleep(delay)
                delay = min(delay * 2, 30)
                continue
            if response.status_code >= 400:
                try:
                    payload = response.json()
                    message = payload.get("message") or payload.get("code") or response.text[:200]
                except ValueError:
                    message = response.text[:200]
                raise EvalApiError(f"HTTP {response.status_code}: {message}")
            try:
                return _unwrap(response.json())
            except ValueError as exc:
                raise EvalApiError(f"响应不是 JSON: {response.text[:200]}") from exc
        raise EvalApiError(f"请求失败（已重试）: {last_error}")

    def ensure_healthy(self) -> None:
        try:
            health = self._client.get("/health")
        except httpx.TransportError as exc:
            raise EvalApiError(f"无法连接识别服务 {self.base_url}: {exc}") from exc
        if health.status_code >= 400:
            raise EvalApiError(f"健康检查失败 HTTP {health.status_code}")
        try:
            ready = self._client.get("/ready")
            payload = ready.json() if ready.headers.get("content-type", "").startswith("application/json") else {}
        except Exception:  # noqa: BLE001
            print("警告: /ready 检查失败，继续尝试识别", file=sys.stderr)
            return
        status = payload.get("status") if isinstance(payload, dict) else None
        if status and status != "ok":
            print(
                f"警告: 服务就绪状态为 {status} "
                f"(redis={payload.get('redis')}, celery={payload.get('celery')})，识别可能失败",
                file=sys.stderr,
            )

    def create_device_session(self, device_id: Optional[str] = None) -> str:
        chosen = (device_id or self.device_id)[:128]
        if len(chosen) < 8:
            chosen = f"{chosen}-eval000"
        data = self._request("POST", "/api/v1/public/sessions/device", json={"device_id": chosen})
        token = data.get("session_token") if isinstance(data, dict) else None
        if not token:
            raise EvalApiError("创建设备会话失败：缺少 session_token")
        self._token = token
        return token

    def create_recognition_session(self) -> str:
        data = self._request("POST", "/api/v1/public/recognize")
        public_id = data.get("public_id") if isinstance(data, dict) else None
        if not public_id:
            raise EvalApiError("创建识别会话失败：缺少 public_id")
        return str(public_id)

    def upload_images(self, public_id: str, paths: Sequence[Path], purpose: str, sort_order: int = 0) -> Any:
        handles = []
        try:
            files = []
            for path in paths:
                handle = path.open("rb")
                handles.append(handle)
                files.append(("files", (path.name, handle, "image/jpeg")))
            return self._request(
                "POST",
                f"/api/v1/public/recognize/{public_id}/images",
                data={"purpose": purpose, "sort_order": str(sort_order)},
                files=files,
            )
        finally:
            for handle in handles:
                handle.close()

    def start(self, public_id: str) -> Any:
        return self._request("POST", f"/api/v1/public/recognize/{public_id}/start")

    def get_session(self, public_id: str) -> Dict[str, Any]:
        data = self._request("GET", f"/api/v1/public/recognize/{public_id}")
        if not isinstance(data, dict):
            raise EvalApiError("识别会话响应格式不正确")
        return data

    def poll_until_done(self, public_id: str) -> Dict[str, Any]:
        deadline = time.monotonic() + self.poll_timeout
        last: Dict[str, Any] = {}
        while time.monotonic() < deadline:
            last = self.get_session(public_id)
            if last.get("status") in {"done", "failed"}:
                return last
            time.sleep(self.poll_interval)
        raise EvalApiError(f"识别超时（{self.poll_timeout:.0f}s），最后状态={last.get('status')}")

    def recognize(self, upload: UploadSet) -> Dict[str, Any]:
        # One independent device session per part so eval does not share the
        # public API's per-device 30 req/min bucket.
        session_id = f"{self.device_id}-{uuid.uuid4().hex[:16]}"[:128]
        self.create_device_session(session_id)
        public_id = self.create_recognition_session()
        self.upload_images(public_id, [upload.overall.path], purpose="part_overall", sort_order=0)
        if upload.others:
            self.upload_images(
                public_id,
                [photo.path for photo in upload.others],
                purpose="part",
                sort_order=1,
            )
        self.start(public_id)
        return self.poll_until_done(public_id)


def default_output_path() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return ROOT / "backend" / "logs" / f"realphoto_eval_{stamp}.xlsx"


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _set_widths(ws: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _bool_label(value: Any) -> str:
    if value is True:
        return "是"
    if value is False:
        return "否"
    return ""


def _join(values: Iterable[str]) -> str:
    return "; ".join(values)


class EvalWorkbook:
    def __init__(self, path: Path) -> None:
        self.path = path
        if path.exists():
            self.wb = load_workbook(path)
        else:
            self.wb = Workbook()
            default = self.wb.active
            if default is not None:
                self.wb.remove(default)
        self._ensure_sheet(RESULT_SHEET, RESULT_HEADERS, [12, 18, 10, 16, 40, 50, 22, 12, 10, 16, 16, 16, 10, 30, 40, 24, 10, 20, 12, 20])
        self._ensure_sheet(SAMPLE_SHEET, SAMPLE_HEADERS, [12, 18, 10, 20, 12, 20, 10])
        if INFO_SHEET not in self.wb.sheetnames:
            ws = self.wb.create_sheet(INFO_SHEET)
            ws.append(["键", "值"])
            _style_header(ws)
            _set_widths(ws, [24, 80])

    def _ensure_sheet(self, name: str, headers: Sequence[str], widths: Sequence[int]) -> Worksheet:
        if name in self.wb.sheetnames:
            return self.wb[name]
        ws = self.wb.create_sheet(name)
        ws.append(list(headers))
        _style_header(ws)
        _set_widths(ws, widths)
        return ws

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.path)

    def successful_material_codes(self) -> Set[str]:
        """Materials whose latest result is status=done (skip on resume)."""
        ws = self.wb[RESULT_SHEET]
        headers = [cell.value for cell in ws[1]]
        try:
            code_idx = headers.index("物料代码")
            status_idx = headers.index("状态")
        except ValueError:
            return set()
        latest: Dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = row[code_idx] if code_idx < len(row) else None
            if not code:
                continue
            status = str(row[status_idx] or "") if status_idx < len(row) else ""
            latest[str(code)] = status
        return {code for code, status in latest.items() if status == "done"}

    def completed_material_codes(self) -> Set[str]:
        return self.successful_material_codes()

    def _row_index_for_material(self, material_code: str) -> Optional[int]:
        ws = self.wb[RESULT_SHEET]
        headers = [cell.value for cell in ws[1]]
        try:
            index = headers.index("物料代码")
        except ValueError:
            return None
        found: Optional[int] = None
        for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            value = row[index] if index < len(row) else None
            if value is not None and str(value) == material_code:
                found = excel_row
        return found

    def write_run_info(self, rows: Sequence[tuple[str, Any]]) -> None:
        ws = self.wb[INFO_SHEET]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for key, value in rows:
            ws.append([key, "" if value is None else str(value)])

    def write_sample_list(
        self,
        groups: Sequence[MaterialGroup],
        uploads: Dict[str, UploadSet],
        completed: Set[str],
    ) -> None:
        ws = self.wb[SAMPLE_SHEET]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for group in groups:
            upload = uploads[group.material_code]
            ws.append(
                [
                    group.prefix,
                    group.material_code,
                    len(group.photos),
                    _join(group.available_angles()),
                    len(upload.all_photos),
                    _join(upload.angles),
                    "是" if group.material_code in completed else "否",
                ]
            )
        ws.auto_filter.ref = ws.dimensions

    def upsert_result(self, row: Dict[str, Any], overall_path: Optional[Path] = None) -> None:
        ws = self.wb[RESULT_SHEET]
        excel_row = self._row_index_for_material(str(row.get("物料代码") or ""))
        if excel_row is None:
            ws.append([row.get(header, "") for header in RESULT_HEADERS])
            excel_row = ws.max_row
        else:
            judge_idx = RESULT_HEADERS.index("人工判定") + 1
            note_idx = RESULT_HEADERS.index("备注") + 1
            old_judge = ws.cell(excel_row, judge_idx).value
            old_note = ws.cell(excel_row, note_idx).value
            for col, header in enumerate(RESULT_HEADERS, start=1):
                ws.cell(excel_row, col).value = row.get(header, "")
            if old_judge:
                ws.cell(excel_row, judge_idx).value = old_judge
            if old_note:
                ws.cell(excel_row, note_idx).value = old_note
        path_cell = ws.cell(excel_row, RESULT_HEADERS.index("图片路径") + 1)
        path_cell.alignment = WRAP
        if overall_path is not None and overall_path.exists():
            path_cell.hyperlink = overall_path.as_uri()
            path_cell.font = LINK_FONT
        ws.cell(excel_row, RESULT_HEADERS.index("候选") + 1).alignment = WRAP
        ws.cell(excel_row, RESULT_HEADERS.index("识别依据") + 1).alignment = WRAP
        ws.auto_filter.ref = ws.dimensions

    def append_result(self, row: Dict[str, Any], overall_path: Optional[Path] = None) -> None:
        self.upsert_result(row, overall_path=overall_path)


def build_result_row(
    *,
    group: MaterialGroup,
    upload: UploadSet,
    session: Optional[Dict[str, Any]] = None,
    error: str = "",
    elapsed_s: float = 0.0,
) -> Dict[str, Any]:
    result = (session or {}).get("result") or {}
    candidates = result.get("candidates") or []
    tested_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return {
        "品类前缀": group.prefix,
        "物料代码": group.material_code,
        "上传张数": len(upload.all_photos),
        "角度编号": _join(upload.angles),
        "图片文件名": _join(upload.filenames),
        "图片路径": _join(str(photo.path) for photo in upload.all_photos),
        "public_id": (session or {}).get("public_id", ""),
        "状态": (session or {}).get("status", "error" if error else ""),
        "是否汽配": _bool_label(result.get("is_auto_part") if session else None),
        "识别品类": result.get("part_name") or "",
        "映射品类": result.get("mapped_category_name") or "",
        "映射品类代码": result.get("mapped_category_code") or "",
        "置信度": result.get("confidence") if result.get("confidence") is not None else "",
        "识别依据": result.get("reason") or "",
        "候选": json.dumps(candidates, ensure_ascii=False) if candidates else "",
        "错误信息": error or (session or {}).get("error_message") or "",
        "耗时秒": round(elapsed_s, 1),
        "测试时间": tested_at,
        "人工判定": "",
        "备注": "",
    }


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按品类前缀抽样实拍图并调用识别 API，结果写入 Excel")
    parser.add_argument("--images-dir", required=True, help="实拍图目录")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help="识别服务地址")
    parser.add_argument("--per-prefix", type=int, default=DEFAULT_PER_PREFIX, help="每种品类前缀抽取的物料数")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED, help="随机种子")
    parser.add_argument("--max-images", type=int, default=DEFAULT_MAX_IMAGES, help="每个物料最多上传张数（含整体外观）")
    parser.add_argument("--limit", type=int, default=None, help="最多识别物料数，用于调试")
    parser.add_argument("--dry-run", action="store_true", help="只抽样并写 Excel，不调用识别")
    parser.add_argument("--resume", type=Path, default=None, help="从已有 xlsx 继续，跳过已成功物料并重试失败项")
    parser.add_argument("--output", type=Path, default=None, help="输出 Excel 路径")
    parser.add_argument("--poll-interval", type=float, default=3.0, help="轮询间隔秒")
    parser.add_argument("--poll-timeout", type=float, default=180.0, help="单次识别超时秒")
    parser.add_argument("--device-id", default=DEFAULT_DEVICE_ID, help="设备 ID 前缀；每个物料会再拼独立后缀")
    return parser.parse_args(argv)


def _print_prefix_table(title: str, counts: Dict[str, int]) -> None:
    print(title)
    total = 0
    for prefix, count in counts.items():
        print(f"  {prefix:<8} {count}")
        total += count
    print(f"  合计     {total}")


def run(args: argparse.Namespace) -> int:
    images_dir = Path(args.images_dir).expanduser().resolve()
    groups = index_photos(images_dir)
    eligible = eligible_groups(groups)
    sampled = sample_by_prefix(eligible, per_prefix=args.per_prefix, seed=args.seed)
    if args.limit is not None:
        sampled = sampled[: max(args.limit, 0)]

    uploads = {group.material_code: select_upload_images(group, max_images=args.max_images) for group in sampled}

    output = Path(args.resume).expanduser().resolve() if args.resume else (args.output or default_output_path())
    output = output.expanduser().resolve()
    workbook = EvalWorkbook(output)
    completed = workbook.successful_material_codes() if args.resume else set()
    pending = [group for group in sampled if group.material_code not in completed]
    started_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def refresh_meta(*, finished: bool = False, failures: int = 0) -> None:
        workbook.write_run_info(
            [
                ("图片目录", images_dir),
                ("API", args.base_url),
                ("seed", args.seed),
                ("per_prefix", args.per_prefix),
                ("max_images", args.max_images),
                ("dry_run", args.dry_run),
                ("limit", args.limit if args.limit is not None else ""),
                ("合格物料总数", len(eligible)),
                ("抽样物料数", len(sampled)),
                ("已测物料数", len(completed)),
                ("待测物料数", len(sampled) - len(completed)),
                ("失败数", failures if finished and not args.dry_run else ""),
                ("开始时间", started_at),
                ("结束时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S") if finished else ""),
                ("合格前缀统计", json.dumps(prefix_counts(eligible), ensure_ascii=False)),
                ("抽样前缀统计", json.dumps(prefix_counts(sampled), ensure_ascii=False)),
            ]
        )
        workbook.write_sample_list(sampled, uploads, completed)
        workbook.save()

    refresh_meta()

    _print_prefix_table("合格物料（按前缀）:", prefix_counts(eligible))
    _print_prefix_table("本次抽样（按前缀）:", prefix_counts(sampled))
    print(f"Excel: {output}")

    if args.dry_run:
        print(f"dry-run：已写入抽样清单，待测 {len(pending)} 个物料")
        return 0

    if not pending:
        print("没有待测物料（可能已全部完成）")
        return 0

    failures = 0
    with RecognizeClient(
        args.base_url,
        poll_interval=args.poll_interval,
        poll_timeout=args.poll_timeout,
        device_id=args.device_id,
    ) as client:
        client.ensure_healthy()
        for index, group in enumerate(pending, start=1):
            upload = uploads[group.material_code]
            print(
                f"[{index}/{len(pending)}] {group.material_code} "
                f"上传 {len(upload.all_photos)} 张 ({_join(upload.angles)})",
                flush=True,
            )
            started = time.monotonic()
            session: Optional[Dict[str, Any]] = None
            error = ""
            try:
                session = client.recognize(upload)
                if session.get("status") == "failed":
                    error = session.get("error_message") or "识别失败"
                    failures += 1
            except Exception as exc:  # noqa: BLE001
                error = str(exc)
                failures += 1
            elapsed = time.monotonic() - started
            row = build_result_row(
                group=group,
                upload=upload,
                session=session,
                error=error,
                elapsed_s=elapsed,
            )
            workbook.upsert_result(row, overall_path=upload.overall.path)
            completed.add(group.material_code)
            refresh_meta(failures=failures)
            part_name = row.get("识别品类") or "-"
            confidence = row.get("置信度")
            print(
                f"    -> 状态={row.get('状态') or 'error'} 品类={part_name} 置信度={confidence} 耗时={elapsed:.1f}s"
                + (f" 错误={error}" if error else ""),
                flush=True,
            )

    refresh_meta(finished=True, failures=failures)
    print(f"完成：成功 {len(pending) - failures}，失败 {failures}，结果 {output}")
    return 1 if failures else 0


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    try:
        return run(args)
    except (FileNotFoundError, ValueError, EvalApiError) as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
